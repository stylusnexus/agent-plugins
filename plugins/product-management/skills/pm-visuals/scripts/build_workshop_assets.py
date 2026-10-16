#!/usr/bin/env python3
"""Build editable workshop artwork. Requires reportlab and python-pptx.

One small scene model feeds SVG, PDF and native PowerPoint shape renderers.
Coordinates and type sizes are points on a 960 x 540 canvas.
"""
import argparse
import csv
import html
import json
import math
from pathlib import Path

from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor
from reportlab.pdfbase.pdfmetrics import stringWidth
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

W, H = 960, 540
INK, MUTED, TEAL = '#203B3B', '#526B69', '#087F78'
PAPER, PALE, LINE, GOLD = '#FAF9F5', '#E6F2EE', '#CADBD5', '#C78936'
DEFAULT = Path(__file__).resolve().parents[1] / 'assets'
SCENES = []


class Scene:
    def __init__(self, slug, title, subtitle, adapted=False):
        self.slug, self.title, self.ops = slug, title, []
        self.adapted = adapted
        self.rect(0, 0, W, H, PAPER)
        self.text(36, 22, 'STYLUS NEXUS  /  PRODUCT WORKSHOP', 9, TEAL, True)
        self.text(36, 48, title, 28, INK, True)
        self.text(36, 88, subtitle, 11, MUTED)
        self.line(36, 117, 924, 117, LINE)
        self.line(36, 498, 924, 498, LINE)
        foot = 'Original Stylus Nexus artwork · MIT · Working template: replace prompts with evidence.'
        if slug == 'strategy-canvas':
            foot = 'Adapted layout: Paweł Huryn / The Product Compass · Product Strategy Canvas v1.2 · CC BY-SA 4.0'
        elif adapted:
            foot = 'Adapted layout: Paweł Huryn / Product Compass + Aatir Abdul Rauf / Behind Product Lines · CC BY-SA 4.0'
        self.text(36, 512, foot, 8, MUTED)
        SCENES.append(self)

    def rect(self, x, y, w, h, fill=PALE, stroke=None):
        self.ops.append(('rect', x, y, w, h, fill, stroke))

    def line(self, x, y, x2, y2, color=LINE, width=1):
        self.ops.append(('line', x, y, x2, y2, color, width))

    def text(self, x, y, txt, size=12, color=INK, bold=False):
        self.ops.append(('text', x, y, txt, size, color, bold))

    def card(self, x, y, w, h, title, lines, number=None):
        self.rect(x, y, w, h, '#FFFFFF', LINE)
        self.rect(x, y, 4, h, TEAL)
        self.text(x+16, y+13, (str(number)+'  ' if number else '')+title, 13, TEAL, True)
        for i, txt in enumerate(lines):
            self.text(x+16, y+39+i*18, txt, 10.5, MUTED)


def pareto(path, unit='Ticket count', context='Fictional support-ticket example', status='fictional'):
    with path.open(newline='', encoding='utf-8') as f:
        rows = list(csv.DictReader(f))
    data = sorted([(r['category'].strip(), float(r['count'])) for r in rows], key=lambda r: -r[1])
    if len(data) != 5 or any(not math.isfinite(v) or v <= 0 for _, v in data):
        raise ValueError('Pareto CSV requires exactly five positive finite counts')
    if any(len(k) > 17 for k, _ in data):
        raise ValueError('Pareto category labels must be at most 17 characters')
    if any(not k for k, _ in data) or len({k for k, _ in data}) != 5:
        raise ValueError('Pareto category labels must be nonempty and unique')
    total = sum(v for _, v in data)
    if not math.isfinite(total) or total <= 0:
        raise ValueError('Pareto total must be finite and positive')
    if not unit.strip() or len(unit) > 24:
        raise ValueError('Pareto unit must contain 1–24 characters')
    if not context.strip() or len(context) > 62:
        raise ValueError('Pareto context must contain 1–62 characters')
    top = max(v for _, v in data)
    status_label = 'Fictional example' if status == 'fictional' else 'User-supplied; not independently verified'
    s = Scene('pareto-analysis', 'Find where the volume concentrates',
              'Pareto analysis · '+context+' · Priority still needs judgment.')
    s.pareto_context, s.pareto_unit, s.pareto_status = context, unit, status
    x0, y0, cw, ch = 78, 164, 562, 250
    s.text(78, 136, unit, 10, MUTED, True)
    s.text(591, 136, 'Cumulative %', 10, MUTED, True)
    for i in range(5):
        y = y0+ch-i*ch/4
        s.line(x0, y, x0+cw, y)
        s.text(43, y-6, f'{top*i/4:g}', 10, MUTED)
        s.text(649, y-6, f'{25*i}%', 10, MUTED)
    s.line(x0, y0+ch, x0+cw, y0+ch, INK)
    guideline = y0+ch*.2
    s.line(x0, guideline, x0+cw, guideline, GOLD, 1.5)
    prev, cumulative = None, 0
    for i, (label, count) in enumerate(data):
        cx = x0+56+i*112
        bh = count/top*ch
        s.rect(cx-30, y0+ch-bh, 60, bh, TEAL if i < 3 else '#91B9AF')
        count_label=f'{count:g}'
        count_width=stringWidth(count_label,'Helvetica-Bold',12)
        short_bar=bh < 26 or count_width > 54
        s.text(cx-count_width/2, y0+ch-bh-18 if short_bar else y0+ch-bh+9,
               count_label, 12, INK if short_bar else '#FFFFFF', True)
        cumulative += count
        cy = y0+ch*(1-cumulative/total)
        if prev:
            s.line(*prev, cx, cy, GOLD, 2.5)
        s.rect(cx-3, cy-3, 6, 6, GOLD)
        s.rect(cx+5, cy-18, 35, 14, PAPER)
        s.text(cx+7, cy-17, f'{100*cumulative/total:g}%', 10, '#98601D', True)
        s.text(cx-44, 426, label, 9, INK)
        prev = (cx, cy)
    s.text(495, guideline+6, '80% reference', 9, '#98601D', True)
    s.text(251, 452, 'Category · sorted by measure', 10, MUTED)
    s.card(710, 153, 214, 135, 'Read the pattern', ['First three categories account', f'for {100*sum(v for _, v in data[:3])/total:.1f}% of total measure.', 'The 80% line is a reference,', 'not a universal law.'])
    s.card(710, 305, 214, 170, 'Before prioritizing', ['Check severity and effort.', 'Check sample + time window.', 'Ask what is preventable.', '', 'Decision: ______________', 'Owner / test: ___________'])
    s.text(78, 473, 'Bars: measure (left) · Line: cumulative % (right) · Total: '+f'{total:g}', 9, MUTED)
    s.text(78, 486, status_label, 8, MUTED)
    return s


def build_scenes(csv_path, pareto_unit='Ticket count', pareto_context='Fictional support-ticket example', pareto_status='fictional'):
    SCENES.clear()
    pareto(csv_path, pareto_unit, pareto_context, pareto_status)
    s = Scene('five-forces', 'Read the pressures around an industry',
              "Porter's Five Forces · Define the industry and geography before judging competitive pressure.")
    for a,b,c,d in [(480,239,480,256),(480,347,480,364),(312,300,336,300),(624,300,648,300)]:
        s.line(a,b,c,d,TEAL,2)
    s.card(336, 137, 288, 102, 'Threat of new entrants', ['Entry barriers, scale, switching costs?', 'Evidence: __________  Pressure: ____'])
    s.card(36, 251, 276, 105, 'Supplier bargaining power', ['Concentration? Unique inputs?', 'Switching costs? Evidence: ______'])
    s.rect(336, 256, 288, 91, TEAL)
    s.text(356, 274, 'Rivalry among existing firms', 15, '#FFFFFF', True)
    s.text(356, 302, 'Growth, differentiation, exit barriers?', 11, '#FFFFFF')
    s.text(356, 323, 'Industry boundary: ______________', 10, '#FFFFFF')
    s.card(648, 251, 276, 105, 'Buyer bargaining power', ['Concentration? Price sensitivity?', 'Alternatives? Evidence: _________'])
    s.card(336, 364, 288, 104, 'Threat of substitutes', ['Other ways to meet the same need?', 'Price / performance? Switching ease?'])
    s.text(40, 382, 'WORKSHOP QUESTION', 9, TEAL, True)
    s.text(40, 402, 'Which pressure changes our choices?', 11)
    s.text(40, 424, 'Response: _____________________', 10, MUTED)
    s.text(650, 391, 'Framework: Michael E. Porter', 10, MUTED)
    s.text(650, 413, 'Assess evidence, direction and timing.', 10, MUTED)
    s.text(650, 435, 'High / medium / low is a discussion aid.', 10, MUTED)
    s = Scene('outcome-roadmap', 'Sequence outcomes, then commit to learning',
              'Now / Next / Later · Horizons express confidence and sequence, not promised delivery dates.')
    for x,title,tag,lines in [(36,'NOW','Act + measure',['Outcome: ___________________','Baseline → target: ___________','Smallest bet: _______________','Owner / review: _____________']), (338,'NEXT','Prepare + validate',['Outcome: ___________________','What must be true? __________','Evidence to advance: ________','Dependency: ________________']), (640,'LATER','Explore + preserve options',['Opportunity: _______________','Why it matters: _____________','Signal to revisit: ___________','What we will not promise: ____'])]:
        s.rect(x,140,284,49,TEAL)
        s.text(x+17,151,title,20,'#FFFFFF',True)
        s.text(x+17,207,tag,12,TEAL,True)
        for i,line in enumerate(lines):
            s.text(x+17,244+i*40,line,12)
        s.line(x+17,420,x+267,420)
    s.text(53,451,'Decision at review: continue / change / stop     ·     Link each bet to an outcome and current evidence.',12,MUTED)
    s = Scene('value-curve', 'Make differences in value explicit',
              'Illustrative strategy profile · Fictional relative scores; validate with your chosen customer segment.')
    labels=['Low effort','Speed','Flexibility','Human help','Low cost']
    for i in range(5):
        y=406-i*56
        s.line(95,y,650,y)
        s.text(70,y-6,str(i+1),10,MUTED)
    s.text(43,139,'Relative offering level · 1 low, 5 high',11,MUTED)
    for values,color in [([2,3,5,2,4],GOLD),([4,4,2,5,2],TEAL)]:
        last=None
        for i,v in enumerate(values):
            x,y=110+i*130,406-(v-1)*56
            if last:
                if color == TEAL:
                    # Short vector segments preserve dash styling across formats.
                    dx,dy=x-last[0],y-last[1]
                    for j in range(0,16,2):
                        s.line(last[0]+dx*j/16,last[1]+dy*j/16,last[0]+dx*(j+1)/16,last[1]+dy*(j+1)/16,color,3)
                else:s.line(*last,x,y,color,3)
            s.rect(x-4,y-4,8,8,color)
            last=(x,y)
    for i,label in enumerate(labels):s.text(82+i*130,425,label,10)
    s.text(215,457,'Customer-valued factors →',11,MUTED)
    s.card(704,145,220,152,'Compare deliberately',['Dashed teal: proposed offer','Solid gold: current alternative','','Segment: ______________','Basis of score: __________'])
    s.card(704,315,220,164,'Choose the trade-offs',['Eliminate: ______________','Reduce: _________________','Raise: __________________','Create: _________________','Test value, not just difference.'])
    s = Scene('assumption-map', 'Test the assumptions that can break the plan',
              'Map one claim per note · Importance means impact if false; evidence means relevant, reliable support.')
    s.rect(96,152,640,276,'#FFFFFF',LINE)
    s.rect(96,152,320,138,PALE)
    s.line(416,152,416,428)
    s.line(96,290,736,290)
    for x,y,t in [(112,166,'TEST FIRST'),(432,166,'MONITOR + PROTECT'),(112,305,'EXPLORE CHEAPLY'),(432,305,'DEFER')]:s.text(x,y,t,11,TEAL,True)
    s.text(96,132,'High importance ↑',10,MUTED)
    s.text(96,440,'Little evidence',10,MUTED)
    s.text(590,440,'Strong evidence →',10,MUTED)
    s.text(36,465,'Low importance at bottom · Move notes only when new evidence warrants it.',11,MUTED)
    s.rect(122,203,264,60,'#F5E8CA')
    s.text(134,214,'H1: [segment] will change behavior',11)
    s.text(134,236,'because [benefit] matters enough.',11)
    s.card(758,152,166,276,'Plan a test',['Claim ID: ________','','Smallest test:','________________','','Decision threshold:','________________','','Owner / date:','________________'])
    s = Scene('customer-journey', 'Follow the job through the customer’s eyes',
              'Define one segment + scenario · Separate observed evidence from hypotheses and team ideas.')
    stages=['Trigger','Explore','Choose','Use','Continue']
    for i,title in enumerate(stages):
        x=166+i*151
        s.rect(x,143,145,43,TEAL)
        s.text(x+10,155,title,14,'#FFFFFF',True)
    rows=[('Customer goal','What progress?'),('Action + touchpoint','What happens?'),('Friction / emotion','Where is the effort?'),('Evidence','Source / unknown'),('Opportunity','Next test / owner')]
    for j,(label,prompt) in enumerate(rows):
        y=194+j*55
        s.text(36,y+11,label,10,TEAL,True)
        for i in range(5):
            x=166+i*151
            s.rect(x,y,145,49,'#FFFFFF',LINE)
            s.text(x+9,y+10,prompt,10,MUTED)
            s.text(x+9,y+29,'________________',10,LINE)
    s.text(36,480,'Scenario: ____________________    Segment: ____________________    Interview / source: ____________________',10,MUTED)
    s = Scene('strategy-canvas', 'Product Strategy Canvas',
              'Product / owner: __________________    Segment: __________________    Date / status: __________________',True)
    fields=[('Vision',['Future worth creating','Values that shape choices']),('Market Segments',['Jobs + chosen context','Who are we excluding?']),('Relative Costs',['Advantage + cost drivers','Broad or focused scope?']),('Value Proposition',['Before → how → after','Alternative + difference']),('Trade-offs',['What we decline + why','Sacrifice + revisit trigger']),('Key Metrics',['Outcome + guardrail','Baseline / target / source']),('Growth',['Channels + growth motion','Retention logic + evidence']),('Capabilities',['Abilities + resources','Current gaps + partners']),('Can’t / Won’t',['Who could copy us?','Barriers + incentive conflicts']),('Ask Yourself',['Do choices reinforce?','Tensions + what must be true'])]
    for i,(title,lines) in enumerate(fields):
        x,y=36+(i%5)*179,139+(i//5)*137
        s.card(x,y,172,125,title,lines)
        s.text(x+16,y+86,'Choice: ______________',9,MUTED)
        s.text(x+16,y+104,'Evidence: ____________',9,MUTED)
    s.rect(36,427,888,54,PALE)
    s.text(49,437,'Consequential choice: __________________   Riskiest assumption / test: __________________',11)
    s.text(49,458,'Owner: __________________    Next review trigger: __________________    Linked objective: __________________',10,MUTED)
    s = Scene('value-proposition', 'Connect customer progress to a credible promise',
              'Product: ______________    Segment / context: ______________    Decision: ______________    Status: draft',True)
    cols=[('01  Job / step',['What progress matters?','When and for whom?']),('02  Existing situation',['What happens today?','Cost / friction / consequence']),('03  Enabling capability',['How does it work?','Why would behavior change?']),('04  Desired benefit',['What becomes better?','Observable outcome'])]
    for i,(title,lines) in enumerate(cols):
        s.card(36+i*225,142,213,139,title,lines)
        s.text(52+i*225,235,'V1: _____________________',10,MUTED)
        s.text(52+i*225,258,'________________________',10,LINE)
    for x,title,lines in [(36,'Alternatives + difference',['Include workarounds and doing nothing.','Why would this segment switch?']), (338,'Evidence + unknowns',['Which link in V1 is supported?','Source / hypothesis / confidence:']), (640,'Next test',['What evidence would change the decision?','Owner / threshold / date:'])]:
        s.card(x,298,284,108,title,lines)
        s.text(x+16,382,'________________________________',10,LINE)
    s.rect(36,424,888,57,PALE)
    s.text(49,434,'Optional message: For [audience], achieve [benefit] through [understandable mechanism].',11)
    s.text(49,456,'Link to V1 evidence. Test comprehension, relevance and credibility; test adoption and outcomes separately.',10,MUTED)


def svg(s, path):
    out=[f'<svg xmlns="http://www.w3.org/2000/svg" width="{W}" height="{H}" viewBox="0 0 {W} {H}">',f'<title>{html.escape(s.title)}</title>']
    for op in s.ops:
        k,*a=op
        if k=='rect':
            x,y,w,h,fill,stroke=a
            out.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" fill="{fill}" stroke="{stroke or "none"}"/>')
        elif k=='line':
            x,y,x2,y2,col,width=a
            out.append(f'<line x1="{x}" y1="{y}" x2="{x2}" y2="{y2}" stroke="{col}" stroke-width="{width}"/>')
        else:
            x,y,t,size,col,bold=a
            out.append(f'<text x="{x}" y="{y+size}" font-family="Arial, Helvetica, sans-serif" font-size="{size}" font-weight="{700 if bold else 400}" fill="{col}">{html.escape(t)}</text>')
    path.write_text('\n'.join(out+['</svg>']),encoding='utf-8')


def pdf_page(s, c):
    for op in s.ops:
        k,*a=op
        if k=='rect':
            x,y,w,h,fill,stroke=a
            c.setFillColor(HexColor(fill))
            if stroke:c.setStrokeColor(HexColor(stroke))
            c.setLineWidth(1)
            c.rect(x,H-y-h,w,h,fill=1,stroke=bool(stroke))
        elif k=='line':
            x,y,x2,y2,col,width=a
            c.setStrokeColor(HexColor(col));c.setLineWidth(width)
            c.line(x,H-y,x2,H-y2)
        else:
            x,y,t,size,col,bold=a
            c.setFont('Helvetica-Bold' if bold else 'Helvetica',size)
            c.setFillColor(HexColor(col))
            # WinAnsi Helvetica lacks these glyphs; keep text readable.
            t=t.replace('→',' > ').replace('↑','^').replace('ł','l')
            c.drawString(x,H-y-size,t)
    c.showPage()


def color(c):return RGBColor.from_string(c.lstrip('#'))


def ppt_page(s, prs):
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    for op in s.ops:
        k,*a=op
        if k=='rect':
            x,y,w,h,fill,stroke=a
            sh=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(x),Pt(y),Pt(w),Pt(h))
            sh.fill.solid();sh.fill.fore_color.rgb=color(fill)
            if stroke:sh.line.color.rgb=color(stroke);sh.line.width=Pt(1)
            else:sh.line.fill.background()
        elif k=='line':
            x,y,x2,y2,col,width=a
            sh=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Pt(x),Pt(y),Pt(x2),Pt(y2))
            sh.line.color.rgb=color(col);sh.line.width=Pt(width)
        else:
            x,y,t,size,col,bold=a
            tw=min(W-x-20,stringWidth(t,'Helvetica-Bold' if bold else 'Helvetica',size)*1.06+8)
            sh=slide.shapes.add_textbox(Pt(x),Pt(y),Pt(tw),Pt(size*1.5+3))
            tf=sh.text_frame;tf.word_wrap=False
            tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
            p=tf.paragraphs[0];p.text=t;p.font.name='Arial';p.font.size=Pt(size)
            p.font.bold=bold;p.font.color.rgb=color(col)
    slide.notes_slide.notes_text_frame.text=json.dumps(metadata(s),ensure_ascii=False,indent=2)


def metadata(s):
    sources=[]
    if s.slug=='pareto-analysis':sources=['https://asq.org/quality-resources/pareto']
    if s.slug=='five-forces':sources=['https://www.isc.hbs.edu/strategy/business-strategy/Pages/the-five-forces.aspx']
    if s.slug=='outcome-roadmap':sources=['https://www.prodpad.com/blog/invented-now-next-later-roadmap/']
    if s.slug=='strategy-canvas':sources=['https://www.productcompass.pm/p/product-strategy-canvas']
    if s.slug=='value-proposition':sources=['https://www.productcompass.pm/p/how-to-design-value-proposition-template']
    return {'id':s.slug,'title':s.title,'license':'CC-BY-SA-4.0' if s.adapted else 'MIT','source_urls':sources,
            'license_url':'https://creativecommons.org/licenses/by-sa/4.0/' if s.adapted else 'https://opensource.org/license/mit/',
            'framework_credit':{'five-forces':'Michael E. Porter','outcome-roadmap':'Janna Bastow (Now / Next / Later)'}.get(s.slug),
            'attribution':('Paweł Huryn / The Product Compass, Product Strategy Canvas v1.2, © 2024' if s.slug=='strategy-canvas' else 'The Product Compass & Behind Product Lines; Paweł Huryn and Aatir Abdul Rauf' if s.adapted else 'Stylus Nexus'),
            **({'pareto_context':s.pareto_context,'pareto_unit':s.pareto_unit,'pareto_status':s.pareto_status} if s.slug=='pareto-analysis' else {}),
            'changes':'Original vector layout adapted from the bundled Markdown template; prompts condensed for workshop use.' if s.adapted else 'Original Stylus Nexus vector artwork; no source diagram reproduced.',
            'outputs':[s.slug+'.svg',s.slug+'.pdf'],'combined_page':SCENES.index(s)+1}


def main():
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--output',type=Path,default=DEFAULT)
    ap.add_argument('--pareto-csv',type=Path,default=DEFAULT/'pareto-fictional-data.csv',help='CSV with category,count headers; exactly five positive finite counts; labels at most 17 characters. Sort and cumulative percentages are recomputed.')
    ap.add_argument('--pareto-unit',default='Ticket count',help='Measure label, 1–24 characters (for example Cost (USD))')
    ap.add_argument('--pareto-context',default='Fictional support-ticket example',help='Dataset context, 1–62 characters')
    ap.add_argument('--pareto-status',choices=['fictional','supplied'],default='fictional',help='Mark fictional example or user-supplied data')
    args=ap.parse_args()
    try:
        build_scenes(args.pareto_csv,args.pareto_unit,args.pareto_context,args.pareto_status)
    except (ValueError,KeyError,TypeError,OSError) as exc:
        ap.error(str(exc))
    args.output.mkdir(parents=True,exist_ok=True)
    prs=Presentation();prs.slide_width=Pt(W);prs.slide_height=Pt(H)
    combined=canvas.Canvas(str(args.output/'workshop-toolkit.pdf'),pagesize=(W,H))
    combined.setTitle('Stylus Nexus product workshop toolkit')
    for s in SCENES:
        svg(s,args.output/(s.slug+'.svg'))
        c=canvas.Canvas(str(args.output/(s.slug+'.pdf')),pagesize=(W,H));c.setTitle(s.title)
        pdf_page(s,c);c.save();pdf_page(s,combined);ppt_page(s,prs)
        (args.output/(s.slug+'.metadata.json')).write_text(json.dumps(metadata(s),indent=2,ensure_ascii=False)+'\n',encoding='utf-8')
    combined.save();prs.save(args.output/'workshop-toolkit.pptx')
    (args.output/'manifest.json').write_text(json.dumps({'title':'Stylus Nexus product workshop toolkit','license_expression':'MIT AND CC-BY-SA-4.0','license_scope':'Pages/slides 1-6: MIT artwork. Pages/slides 7-8: CC-BY-SA-4.0 adaptations; retain attribution and share alike when adapting those works.','size_points':[W,H],'assets':[metadata(s) for s in SCENES]},indent=2,ensure_ascii=False)+'\n',encoding='utf-8')
    print(f'Built {len(SCENES)} SVGs, {len(SCENES)} PDFs, an 8-page PDF and an 8-slide editable PPTX in {args.output}')


if __name__=='__main__':main()
